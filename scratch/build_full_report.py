import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def build_docx_report():
    doc = docx.Document()

    # Page Setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles & Colors
    NAVY = RGBColor(27, 54, 93)      # #1B365D
    SLATE = RGBColor(44, 82, 130)    # #2C5282
    TEAL = RGBColor(43, 108, 176)    # #2B6CB0
    DARK_TEXT = RGBColor(45, 55, 72) # #2D3748

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, color="CBD5E0", sz="4"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_h1(text):
        h = doc.add_heading(text, level=1)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        run = h.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = NAVY
        return h

    def add_h2(text):
        h = doc.add_heading(text, level=2)
        h.paragraph_format.keep_with_next = True
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        run = h.runs[0]
        run.font.name = "Calibri"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = SLATE
        return h

    def add_body(text, bold_prefix=""):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r1 = p.add_run(bold_prefix)
            r1.font.name = "Calibri"
            r1.font.size = Pt(10.5)
            r1.font.bold = True
            r1.font.color.rgb = DARK_TEXT
        r2 = p.add_run(text)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = DARK_TEXT
        return p

    def add_code(code_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F7FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:left w:val="single" w:sz="12" w:space="0" w:color="2B6CB0"/>\n'
            f'  <w:top w:val="none"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:bottom w:val="none"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(0)
        r = cp.add_run(code_text.strip())
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        r.font.color.rgb = DARK_TEXT
        
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    def add_placeholder(caption):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "EDF2F7")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="A0AEC0"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(0)
        
        r1 = cp.add_run("📷 ")
        r1.font.size = Pt(11)
        r2 = cp.add_run(f"[{caption}]")
        r2.font.name = "Calibri"
        r2.font.size = Pt(9.5)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(113, 128, 150)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    def add_table(headers, data):
        tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="CBD5E0", sz="4")

        hdr_cells = tbl.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "1B365D")
            set_cell_margins(hdr_cells[i], top=90, bottom=90, left=100, right=100)
            p = hdr_cells[i].paragraphs[0]
            for run in p.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, row_data in enumerate(data):
            row_cells = tbl.rows[row_idx + 1].cells
            bg_color = "F7FAFC" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=100, right=100)
                p = row_cells[col_idx].paragraphs[0]
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = DARK_TEXT

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1. COVER PAGE
    # -------------------------------------------------------------
    p_title_sub = doc.add_paragraph()
    p_title_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title_sub.paragraph_format.space_before = Pt(36)
    r = p_title_sub.add_run("Web Services & SOA Laboratory")
    r.font.name = "Calibri"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = SLATE

    p_title_main = doc.add_paragraph()
    p_title_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title_main.paragraph_format.space_before = Pt(12)
    p_title_main.paragraph_format.space_after = Pt(24)
    r = p_title_main.add_run("Lab 6 – Docker & Microservices: Decomposing and Running the Backend as Independent Services")
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = NAVY

    # Cover Box
    tbl_cover = doc.add_table(rows=5, cols=2)
    tbl_cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_cover, color="CBD5E0", sz="6")
    
    meta_info = [
        ("Project Title", "CampusConnect Backend Microservices"),
        ("Student Name", "Rishika Shah"),
        ("Enrollment Number", "____________________"),
        ("Division / Batch", "____________________"),
        ("Submission Date", "October 2026")
    ]
    for idx, (k, v) in enumerate(meta_info):
        cell_k = tbl_cover.cell(idx, 0)
        cell_v = tbl_cover.cell(idx, 1)
        cell_k.text = k
        cell_v.text = v
        set_cell_background(cell_k, "EDF2F7")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, top=100, bottom=100, left=150, right=150)
        set_cell_margins(cell_v, top=100, bottom=100, left=150, right=150)
        p_k = cell_k.paragraphs[0]
        p_v = cell_v.paragraphs[0]
        p_k.runs[0].font.bold = True
        p_k.runs[0].font.size = Pt(10)
        p_v.runs[0].font.size = Pt(10)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. OBJECTIVE
    # -------------------------------------------------------------
    add_h1("1. Objective")
    add_body("The primary objectives of Laboratory 6 are as follows:")
    add_body("Decompose the existing monolithic Node.js/Express backend (CampusConnect) into three independently runnable microservices: User Service, Product Service, and Order Service.", "• Service Decomposition: ")
    add_body("Create dedicated Dockerfiles for each microservice to define containerized runtime environments using standard Node base images.", "• Containerization: ")
    add_body("Connect all microservice and database containers using a custom Docker bridge network (campus-network).", "• Docker Networking: ")
    add_body("Configure inter-service communication URLs dynamically using environment variables rather than hardcoded endpoints.", "• Externalized Configuration: ")
    add_body("Orchestrate the entire multi-service application (3 microservices + 3 isolated database instances) using Docker Compose.", "• Multi-Container Orchestration: ")
    add_body("Implement robust inter-service REST communication, validating dependent resources and returning structured HTTP status codes (200, 201, 404, 503).", "• Inter-Service API Verification: ")

    # -------------------------------------------------------------
    # 3. INTRODUCTION / LAB CONTEXT
    # -------------------------------------------------------------
    add_h1("2. Introduction / Lab Context")
    add_body("In previous laboratory assignments (Lab 4 & Lab 5), the CampusConnect backend was developed as a monolithic Express.js REST API communicating with a single MongoDB instance. While containerized, all API routes (students, catalog, orders) shared a single execution thread, codebase, and database connection.")
    add_body("In Lab 6, we transition from monolithic containerization to a Service-Oriented Microservices Architecture (SOA). The application backend is split into three decoupled services: User Service (handling user accounts), Product Service (handling merchandise catalog), and Order Service (handling order placement). Order Service relies on REST APIs exposed by User and Product services for validation, demonstrating true inter-service HTTP communication over a custom Docker network.")

    # -------------------------------------------------------------
    # 4. MICROSERVICES ARCHITECTURE
    # -------------------------------------------------------------
    add_h1("3. Microservices Architecture")
    add_body("The CampusConnect system architecture consists of three independently deployable Node.js microservices and three dedicated MongoDB database containers, connected via the campus-network Docker bridge network.")
    
    add_h2("System Architecture Topology")
    add_code(
"┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
"│                                Client / Postman Consumer                               │\n"
"└──────────────────────────┬──────────────────┬──────────────────┬───────────────────────┘\n"
"                           │ HTTP             │ HTTP             │ HTTP\n"
"                           │ :3001            │ :3002            │ :3003\n"
"                           ▼                  ▼                  ▼\n"
"┌────────────────────────────────────────────────────────────────────────────────────────┐\n"
"│ Docker Host (Docker Network: campus-network)                                           │\n"
"│                                                                                        │\n"
"│   ┌───────────────────────┐   HTTP GET /users/{id}    ┌──────────────────────────┐    │\n"
"│   │     User Service      │◄──────────────────────────┤      Order Service       │    │\n"
"│   │  (user-service:3001)  │ http://user-service:3001  │  (order-service:3003)    │    │\n"
"│   └───────────┬───────────┘                           └────────────┬─────────────┘    │\n"
"│               │                                                    │                  │\n"
"│               │ MongoDB                                            │ HTTP GET /products\n"
"│               ▼                                                    │ http://product-..\n"
"│   ┌───────────────────────┐                           ┌────────────▼─────────────┐    │\n"
"│   │     User Database     │                           │     Product Service      │    │\n"
"│   │   (user-db:27017)     │                           │ (product-service:3002)   │    │\n"
"│   └───────────────────────┘                           └────────────┬─────────────┘    │\n"
"│                                                                    │                  │\n"
"│   ┌───────────────────────┐                           ┌────────────▼─────────────┐    │\n"
"│   │    Order Database     │                           │     Product Database     │    │\n"
"│   │   (order-db:27017)    │                           │   (product-db:27017)     │    │\n"
"│   └───────────────────────┘                           └──────────────────────────┘    │\n"
"└────────────────────────────────────────────────────────────────────────────────────────┘"
    )
    add_body("Note: An API Gateway was not required for this assignment and is conceptually represented by direct port exposure (3001, 3002, 3003) on the host machine.")

    # -------------------------------------------------------------
    # 5. SERVICE DECOMPOSITION
    # -------------------------------------------------------------
    add_h1("4. Service Decomposition")
    add_body("The backend functionality has been partitioned according to business domain boundaries:")
    
    headers_decomp = ["Service Name", "Responsibility", "Port", "Main Resource", "REST Endpoints"]
    data_decomp = [
        ["User Service", "User profile CRUD & department management", "3001", "Users", "GET /users\nGET /users/:id\nPOST /users\nPUT /users/:id\nDELETE /users/:id"],
        ["Product Service", "Product catalog, pricing & stock inventory", "3002", "Products", "GET /products\nGET /products/:id\nPOST /products\nPUT /products/:id\nDELETE /products/:id"],
        ["Order Service", "Order placement & inter-service verification", "3003", "Orders", "GET /orders\nGET /orders/:id\nPOST /orders"]
    ]
    add_table(headers_decomp, data_decomp)

    # -------------------------------------------------------------
    # 6. PROJECT STRUCTURE
    # -------------------------------------------------------------
    add_h1("5. Project Structure")
    add_body("The actual file system structure under D:\\Assignments\\student-api reflects a clean microservices organization:")
    add_code(
"student-api/\n"
"├── user-service/\n"
"│   ├── server.js               # Express REST API (port 3001)\n"
"│   ├── package.json            # Service dependencies (express, mongoose, cors)\n"
"│   ├── Dockerfile              # Docker build configuration\n"
"│   ├── .env                    # Environment variables\n"
"│   └── models/\n"
"│       └── User.js             # Mongoose User model schema\n"
"├── product-service/\n"
"│   ├── server.js               # Express REST API (port 3002)\n"
"│   ├── package.json            # Service dependencies\n"
"│   ├── Dockerfile              # Docker build configuration\n"
"│   ├── .env                    # Environment variables\n"
"│   └── models/\n"
"│       └── Product.js          # Mongoose Product model schema\n"
"├── order-service/\n"
"│   ├── server.js               # Express REST API (port 3003)\n"
"│   ├── package.json            # Service dependencies\n"
"│   ├── Dockerfile              # Docker build configuration\n"
"│   ├── .env                    # Environment variables\n"
"│   └── models/\n"
"│       └── Order.js            # Mongoose Order model schema\n"
"├── compose.yaml                # Master Docker Compose specification\n"
"├── docker-compose.yml          # Mirror Compose configuration file\n"
"├── Microservices-Lab6.postman_collection.json # Verified Postman Collection\n"
"└── README.md                   # System documentation & execution guide"
    )
    add_placeholder("Figure 1: Project Directory Structure – Screenshot to be added")

    # -------------------------------------------------------------
    # 7. USER SERVICE
    # -------------------------------------------------------------
    add_h1("6. User Service")
    add_body("User Service manages user accounts and profiles. It operates independently on port 3001 and connects to user-db (MongoDB).", "• Purpose & Port: ")
    add_body("GET /users, GET /users/:id, POST /users, PUT /users/:id, DELETE /users/:id.", "• Supported REST Endpoints: ")
    add_body("User data is stored exclusively in the user_db database on user-db:27017.", "• Data Ownership: ")
    add_body("Runs as a Node.js 20 container using node server.js.", "• Execution: ")
    add_placeholder("Figure 2: User Service Dockerfile – Screenshot to be added")
    add_placeholder("Figure 3: User API Testing – Screenshot to be added")

    # -------------------------------------------------------------
    # 8. PRODUCT SERVICE
    # -------------------------------------------------------------
    add_h1("7. Product Service")
    add_body("Product Service manages merchandise catalog items, pricing, and stock levels on port 3002, connected to product-db (MongoDB).", "• Purpose & Port: ")
    add_body("GET /products, GET /products/:id, POST /products, PUT /products/:id, DELETE /products/:id.", "• Supported REST Endpoints: ")
    add_body("Product data is stored exclusively in the product_db database on product-db:27017.", "• Data Ownership: ")
    add_placeholder("Figure 4: Product Service Dockerfile – Screenshot to be added")
    add_placeholder("Figure 5: Product API Testing – Screenshot to be added")

    # -------------------------------------------------------------
    # 9. ORDER SERVICE
    # -------------------------------------------------------------
    add_h1("8. Order Service")
    add_body("Order Service coordinates order placement on port 3003 and stores order history in order-db (MongoDB).", "• Purpose & Port: ")
    add_body("GET /orders, GET /orders/:id, POST /orders.", "• Supported REST Endpoints: ")
    add_body("When POST /orders is called with { userId, productId, quantity }:", "• Order Creation Flow: ")
    add_body("Order Service issues HTTP GET to ${USER_SERVICE_URL}/users/${userId} (http://user-service:3001/users/{id}). If HTTP 404 is returned, Order Service halts and returns HTTP 404 User Not Found.", "  1. User Verification: ")
    add_body("Order Service issues HTTP GET to ${PRODUCT_SERVICE_URL}/products/${productId} (http://product-service:3002/products/{id}). If HTTP 404 is returned, Order Service halts and returns HTTP 404 Product Not Found.", "  2. Product Verification: ")
    add_body("Order Service calculates totalPrice = product.price * quantity and saves the order to order-db, returning HTTP 201 Created.", "  3. Order Saving: ")
    add_body("If either User Service or Product Service is offline/unreachable, Order Service catches the fetch error and returns HTTP 503 Service Unavailable.", "  4. Resilience: ")

    # -------------------------------------------------------------
    # 10. SERVICE-TO-SERVICE COMMUNICATION
    # -------------------------------------------------------------
    add_h1("9. Service-to-Service Communication")
    add_body("Order Service communicates strictly over HTTP REST APIs using container DNS names over the shared Docker network:")
    add_code(
"Order Service (:3003)\n"
"    │\n"
"    ├─── GET http://user-service:3001/users/{userId} ───────► User Service (:3001)\n"
"    │\n"
"    └─── GET http://product-service:3002/products/{productId} ─► Product Service (:3002)"
    )
    
    headers_comm = ["Calling Service", "Target Service", "HTTP Method", "Target Endpoint", "Request Data", "Success Result", "Error Handling"]
    data_comm = [
        ["Order Service", "User Service", "GET", "http://user-service:3001/users/{id}", "User ID (path param)", "200 OK + User JSON", "404 if User missing\n503 if Service offline"],
        ["Order Service", "Product Service", "GET", "http://product-service:3002/products/{id}", "Product ID (path param)", "200 OK + Product JSON", "404 if Product missing\n503 if Service offline"]
    ]
    add_table(headers_comm, data_comm)

    # -------------------------------------------------------------
    # 11. DOCKERFILES
    # -------------------------------------------------------------
    add_h1("10. Dockerfiles")
    add_body("Each microservice features a standardized Dockerfile built from node:20 base image:")
    
    headers_df = ["Service Name", "Dockerfile Path", "Base Image", "Exposed Port", "Start Command"]
    data_df = [
        ["User Service", "user-service/Dockerfile", "node:20", "3001", "node server.js"],
        ["Product Service", "product-service/Dockerfile", "node:20", "3002", "node server.js"],
        ["Order Service", "order-service/Dockerfile", "node:20", "3003", "node server.js"]
    ]
    add_table(headers_df, data_df)

    # -------------------------------------------------------------
    # 12. DOCKER NETWORK
    # -------------------------------------------------------------
    add_h1("11. Docker Network")
    add_body("A custom bridge network named campus-network is defined in compose.yaml. All 6 containers (3 microservices + 3 databases) are attached to campus-network. Docker provides automatic embedded DNS resolution, allowing order-service to resolve http://user-service:3001 and http://product-service:3002 without using localhost or hardcoded IP addresses.")

    # -------------------------------------------------------------
    # 13. ENVIRONMENT VARIABLES
    # -------------------------------------------------------------
    add_h1("12. Environment Variables")
    add_body("Configuration values are externalized using environment variables:")
    
    headers_env = ["Variable Name", "Target Service", "Purpose", "Configured Value in Compose"]
    data_env = [
        ["PORT", "All Services", "Service listening port", "3001 / 3002 / 3003"],
        ["MONGODB_URI", "user-service", "User DB connection string", "mongodb://user-db:27017/user_db"],
        ["MONGODB_URI", "product-service", "Product DB connection string", "mongodb://product-db:27017/product_db"],
        ["MONGODB_URI", "order-service", "Order DB connection string", "mongodb://order-db:27017/order_db"],
        ["USER_SERVICE_URL", "order-service", "User Service API endpoint", "http://user-service:3001"],
        ["PRODUCT_SERVICE_URL", "order-service", "Product Service API endpoint", "http://product-service:3002"]
    ]
    add_table(headers_env, data_env)

    # -------------------------------------------------------------
    # 14. DOCKER COMPOSE
    # -------------------------------------------------------------
    add_h1("13. Docker Compose")
    add_body("Multi-container orchestration is defined in compose.yaml and docker-compose.yml:")
    
    headers_cmp = ["Service Container", "Image / Build Source", "Host Port Mapping", "Network"]
    data_cmp = [
        ["user-db", "mongo:latest", "27017:27017", "campus-network"],
        ["user-service", "build: ./user-service", "3001:3001", "campus-network"],
        ["product-db", "mongo:latest", "27018:27017", "campus-network"],
        ["product-service", "build: ./product-service", "3002:3002", "campus-network"],
        ["order-db", "mongo:latest", "27019:27017", "campus-network"],
        ["order-service", "build: ./order-service", "3003:3003", "campus-network"]
    ]
    add_table(headers_cmp, data_cmp)
    
    add_body("Primary Execution Commands:")
    add_code(
"# Start complete application\n"
"docker compose up --build -d\n\n"
"# Verify running containers\n"
"docker compose ps\n\n"
"# Inspect application logs\n"
"docker compose logs"
    )

    # -------------------------------------------------------------
    # 15. DOCKER IMAGE BUILD
    # -------------------------------------------------------------
    add_h1("14. Docker Image Build")
    add_body("Images can be built individually or automatically via Docker Compose:")
    add_code(
"docker build -t student-api-user-service ./user-service\n"
"docker build -t student-api-product-service ./product-service\n"
"docker build -t student-api-order-service ./order-service"
    )
    add_placeholder("Figure 6: Docker Images List – Screenshot to be added")

    # -------------------------------------------------------------
    # 16. RUNNING THE APPLICATION
    # -------------------------------------------------------------
    add_h1("15. Running the Application")
    add_body("Executing docker compose up --build -d launches all 6 containers. Running docker compose ps yields the verified active status:")
    add_code(
"NAME              IMAGE                         STATUS         PORTS\n"
"user-db           mongo:latest                  Up             0.0.0.0:27017->27017/tcp\n"
"user-service      student-api-user-service      Up             0.0.0.0:3001->3001/tcp\n"
"product-db        mongo:latest                  Up             0.0.0.0:27018->27017/tcp\n"
"product-service   student-api-product-service   Up             0.0.0.0:3002->3002/tcp\n"
"order-db          mongo:latest                  Up             0.0.0.0:27019->27017/tcp\n"
"order-service     student-api-order-service     Up             0.0.0.0:3003->3003/tcp"
    )

    # -------------------------------------------------------------
    # 17. POSTMAN TESTING
    # -------------------------------------------------------------
    add_h1("16. Postman Testing")
    add_body("The REST APIs were verified using Microservices-Lab6.postman_collection.json:")
    
    headers_postman = ["Test Case", "Method", "Endpoint", "Expected Result", "Actual Verified Result"]
    data_postman = [
        ["1. Create User", "POST", "http://localhost:3001/users", "201 Created", "201 Created (User ID 101)"],
        ["2. Fetch User", "GET", "http://localhost:3001/users/101", "200 OK", "200 OK (Alice Smith)"],
        ["3. Create Product", "POST", "http://localhost:3002/products", "201 Created", "201 Created (Product ID 501)"],
        ["4. Fetch Product", "GET", "http://localhost:3002/products/501", "200 OK", "200 OK (Campus Hoodie, $45)"],
        ["5. Create Order", "POST", "http://localhost:3003/orders", "201 Created", "201 Created (Total $90, Qty 2)"],
        ["6. Fetch Orders", "GET", "http://localhost:3003/orders", "200 OK", "200 OK (Array of orders)"],
        ["7. Invalid User ID", "POST", "http://localhost:3003/orders", "404 Not Found", "404 Not Found (User '99999' missing)"],
        ["8. Dependency Down", "POST", "http://localhost:3003/orders", "503 Unavailable", "503 Service Unavailable"],
        ["9. Service Recovery", "POST", "http://localhost:3003/orders", "201 Created", "201 Created (Order 1002)"]
    ]
    add_table(headers_postman, data_postman)

    # -------------------------------------------------------------
    # 18. INVALID RESOURCE / 404 TEST
    # -------------------------------------------------------------
    add_h1("17. Invalid Resource / 404 Test")
    add_body("When POST /orders is called with an unrecorded userId (e.g. 99999), Order Service queries User Service, receives HTTP 404, and returns HTTP 404 Not Found with message: \"User with ID '99999' not found.\"")
    add_placeholder("Figure 7: Invalid Resource ID 404 Response – Screenshot to be added")

    # -------------------------------------------------------------
    # 19. DEPENDENCY FAILURE / 503 TEST & RECOVERY
    # -------------------------------------------------------------
    add_h1("18. Dependency Failure / 503 Test & Recovery")
    add_body("1. User Service container stopped: docker stop user-service", "• Execution Step 1: ")
    add_body("Order creation attempted via POST http://localhost:3003/orders.", "• Execution Step 2: ")
    add_body("Order Service attempts to contact http://user-service:3001/users/101, catches connection failure, and returns HTTP 503 Service Unavailable.", "• Result 1: ")
    add_placeholder("Figure 8: Dependency Unavailable 503 Response – Screenshot to be added")
    add_body("User Service restarted: docker start user-service", "• Execution Step 3: ")
    add_body("Order creation retried. Order Service successfully contacts User Service and returns HTTP 201 Created.", "• Result 2: ")
    add_placeholder("Figure 9: Service Recovery Success – Screenshot to be added")

    # -------------------------------------------------------------
    # 20. DATABASE PER SERVICE
    # -------------------------------------------------------------
    add_h1("19. Database Per Service")
    add_body("The application strictly implements the Database-per-Service architectural pattern:")
    add_body("User Service owns user_db on container user-db.", "• User Data Isolation: ")
    add_body("Product Service owns product_db on container product-db.", "• Product Data Isolation: ")
    add_body("Order Service owns order_db on container order-db.", "• Order Data Isolation: ")
    add_body("No service is permitted to directly query or write to another service's MongoDB instance. All cross-domain data exchange takes place over HTTP REST APIs.")

    # -------------------------------------------------------------
    # 21. API & COMMUNICATION EXERCISE MATRIX
    # -------------------------------------------------------------
    add_h1("20. API & Communication Exercise Matrix")
    
    headers_ex = ["Field", "User Service", "Product Service", "Order Service"]
    data_ex = [
        ["Responsibility", "User profile CRUD", "Product catalog management", "Order creation & verification"],
        ["Port", "3001", "3002", "3003"],
        ["Main Resource", "Users", "Products", "Orders"],
        ["Key Endpoints", "GET/POST/PUT/DELETE /users", "GET/POST/PUT/DELETE /products", "POST/GET /orders"],
        ["Data / Database", "user-db (user_db)", "product-db (product_db)", "order-db (order_db)"],
        ["Service-to-Service Call", "None (Target Service)", "None (Target Service)", "Calls User & Product APIs"]
    ]
    add_table(headers_ex, data_ex)

    add_body("Inter-Service Call Summary:")
    add_body("Order Service", "• Calling Service: ")
    add_body("User Service (:3001) / Product Service (:3002)", "• Target Services: ")
    add_body("GET", "• HTTP Method: ")
    add_body("http://user-service:3001/users/{id} & http://product-service:3002/products/{id}", "• Endpoints Called: ")
    add_body("200 OK + Resource JSON", "• Expected Success Response: ")
    add_body("404 Not Found", "• Invalid Resource Response: ")
    add_body("503 Service Unavailable", "• Dependency Failure Response: ")

    # -------------------------------------------------------------
    # 22. EVIDENCE / SCREENSHOTS PLACEHOLDERS SUMMARY
    # -------------------------------------------------------------
    add_h1("21. Evidence / Screenshots Summary")
    add_body("The report incorporates the following figure placeholders ready for screenshot insertion:")
    
    headers_fig = ["Figure #", "Caption Description", "Status"]
    data_fig = [
        ["Figure 1", "Project Directory Structure", "Placeholder ready"],
        ["Figure 2", "User Service Dockerfile", "Placeholder ready"],
        ["Figure 3", "User API Testing (Postman)", "Placeholder ready"],
        ["Figure 4", "Product Service Dockerfile", "Placeholder ready"],
        ["Figure 5", "Product API Testing (Postman)", "Placeholder ready"],
        ["Figure 6", "Docker Images List (docker images)", "Placeholder ready"],
        ["Figure 7", "Invalid Resource ID 404 Response", "Placeholder ready"],
        ["Figure 8", "Dependency Unavailable 503 Response", "Placeholder ready"],
        ["Figure 9", "Service Recovery Success Response", "Placeholder ready"]
    ]
    add_table(headers_fig, data_fig)

    # -------------------------------------------------------------
    # 23. TROUBLESHOOTING
    # -------------------------------------------------------------
    add_h1("22. Troubleshooting & Resolved Issues")
    add_body("During automated testing, PowerShell inline quotes resulted in malformed JSON payload errors. Resolved by executing tests via a Node.js test script using native fetch.", "• PowerShell JSON String Escaping: ")
    add_body("To prevent Order Service from hanging when User or Product Service is offline, AbortSignal.timeout(4000) was integrated into outbound fetch calls, ensuring immediate 503 responses.", "• Outbound Request Timeout: ")

    # -------------------------------------------------------------
    # 24. CONCLUSION
    # -------------------------------------------------------------
    add_h1("23. Conclusion")
    add_body("Lab 6 successfully refactored the monolithic CampusConnect application into three independently runnable, containerized microservices (User, Product, and Order services). Each service owns its database instance and runs inside a dedicated Docker container. Inter-service REST communication over the campus-network bridge network was verified, demonstrating strict domain isolation, controlled error responses (404, 503), and rapid system recovery upon service restart.")

    # -------------------------------------------------------------
    # 25. VIVA QUESTIONS & ANSWERS
    # -------------------------------------------------------------
    add_h1("24. Viva Questions & Answers")
    
    viva_qas = [
        ("1. What is a microservice?", "A microservice is a small, independently deployable software component that handles a single specific business capability and communicates over standard network protocols like HTTP REST."),
        ("2. Why did we divide the backend into three services?", "To achieve loose coupling, independent scalability, isolated fault domains, and distinct database ownership for Users, Products, and Orders."),
        ("3. What is Docker?", "Docker is a containerization platform that packages application code along with all its dependencies into portable, isolated containers."),
        ("4. Why does each service have its own Dockerfile?", "Because each microservice has its own codebase, dependencies, port exposure, and entry point script."),
        ("5. What is Docker Compose?", "Docker Compose is a tool for defining and running multi-container Docker applications using a single YAML configuration file."),
        ("6. What is a Docker network?", "A Docker network is an isolated virtual network bridge that allows containers on the same host to communicate securely."),
        ("7. Why should we not use localhost between containers?", "Inside a container, 'localhost' points strictly to that container itself, not other containers on the host."),
        ("8. What is service-name communication?", "Using Docker Compose container service names (e.g. user-service) as DNS hostnames for inter-container HTTP calls."),
        ("9. Why does Order Service call User Service?", "To verify that the referenced user exists and retrieve user profile details before saving an order."),
        ("10. Why does Order Service call Product Service?", "To verify product existence, check availability, and retrieve product pricing to calculate the order total."),
        ("11. What is database-per-service?", "A pattern where each microservice exclusively owns its database instance, preventing direct database access from other services."),
        ("12. What happens when User Service is unavailable?", "Order Service catches the connection failure and returns HTTP 503 Service Unavailable."),
        ("13. What is HTTP 404?", "An HTTP status code indicating that the requested resource (User or Product ID) was not found."),
        ("14. What is HTTP 503?", "An HTTP status code indicating that a required upstream service is currently unavailable."),
        ("15. What is the role of environment variables?", "To externalize configuration settings like ports and service URLs so code does not need re-compilation between environments."),
        ("16. What is an API?", "Application Programming Interface—a set of defined rules allowing different applications to communicate."),
        ("17. What is REST?", "Representational State Transfer—an architectural style using standard HTTP methods (GET, POST, PUT, DELETE) for stateless web APIs."),
        ("18. What is containerization?", "Packaging an application and its runtime environment into a container image that runs consistently anywhere."),
        ("19. Difference between monolith and microservices?", "A monolith packages all features in a single binary/process; microservices split features into separate, independent services."),
        ("20. How do you start the complete application?", "Run 'docker compose up --build -d' from the project directory.")
    ]

    for q, a in viva_qas:
        add_body(a, f"{q}\nAnswer: ")

    # Save DOCX
    docx_path = r"D:\Assignments\student-api\CampusConnect_Lab6_Report.docx"
    doc.save(docx_path)
    print(f"DOCX report saved successfully at: {docx_path}")
    return docx_path

if __name__ == "__main__":
    build_docx_report()
